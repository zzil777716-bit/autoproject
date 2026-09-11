"""
========================================================================================
📅 [WORKER: THEME & DESKTOP CALENDAR WORKER]
Executes HTS 4-part data collection, updates monthly Excel sheet, and builds Desktop Calendar HTML.
Permanently integrates the 5th tab: [🎯 윗꼬리 눌림목 TOP 30 승률 & 타점 감시]
========================================================================================
"""

import os
import sys
import json
from datetime import datetime
from typing import Dict, Any, List, Optional
try:
    sys.stdout.reconfigure(encoding='utf-8')
    sys.stderr.reconfigure(encoding='utf-8')
except Exception:
    pass

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from collectors.hts_theme_collector import HTSThemeCollector
from storage.gdrive_sync import GDriveSync

class ThemeCalendarWorker:
    def __init__(self, base_dir: str = r"C:\Antigravity"):
        self.base_dir = base_dir
        self.themes_dir = os.path.join(base_dir, "data", "themes")
        self.calendar_dir = os.path.join(base_dir, "data", "calendar")
        self.verification_dir = os.path.join(base_dir, "data", "사용자_검증")
        os.makedirs(self.themes_dir, exist_ok=True)
        os.makedirs(self.calendar_dir, exist_ok=True)
        os.makedirs(self.verification_dir, exist_ok=True)
        
        self.history_file = os.path.join(self.calendar_dir, "market_calendar_history.json")
        self.desktop_html_path = r"C:\Users\HONG\Desktop\증시_주도테마_캘린더.html"
        self.gdrive_html_path = r"G:\내 드라이브\Antigravity\캘린더\증시_주도테마_캘린더.html"
        self.gdrive_sync = GDriveSync()

    def run(self, api_instance=None):
        collector = HTSThemeCollector(api_instance)
        market_data = collector.collect_all()
        now = datetime.now()
        
        # 1. 1달 단위 통합 엑셀 워크북(HTS_테마_주도주_분석_YYYY-MM.xlsx) 갱신
        month_str = now.strftime('%Y-%m')
        date_str = now.strftime('%Y-%m-%d')
        monthly_excel_path = os.path.join(self.themes_dir, f"HTS_테마_주도주_분석_{month_str}.xlsx")
        
        self._update_monthly_excel(monthly_excel_path, date_str, market_data)

        # 2. 캘린더 히스토리 적재
        history = self._load_history()
        history[date_str] = market_data
        self._save_history(history)

        # 3. 바탕화면 캘린더 빌드 (영구 5번째 메뉴 TOP 30 통합)
        self._build_desktop_calendar(history, now.year, now.month)

        # 4. 구글 드라이브 일괄 동기화
        self.gdrive_sync.sync_all(self.base_dir)

    def _update_monthly_excel(self, excel_path: str, date_str: str, market_data: Dict[str, Any]):
        if os.path.exists(excel_path):
            try: wb = openpyxl.load_workbook(excel_path)
            except Exception: wb = openpyxl.Workbook()
        else:
            wb = openpyxl.Workbook()

        if date_str in wb.sheetnames:
            del wb[date_str]

        ws = wb.create_sheet(title=date_str)
        ws.views.sheetView[0].showGridLines = True
        if "Sheet" in wb.sheetnames and len(wb.sheetnames) > 1:
            del wb["Sheet"]

        navy_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
        theme_hdr_fill = PatternFill(start_color="2B6CB0", end_color="2B6CB0", fill_type="solid")
        sub_hdr_fill = PatternFill(start_color="EDF2F7", end_color="EDF2F7", fill_type="solid")
        amber_hdr_fill = PatternFill(start_color="78350F", end_color="78350F", fill_type="solid")
        emerald_hdr_fill = PatternFill(start_color="064E3B", end_color="064E3B", fill_type="solid")
        blue_hdr_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")

        title_font = Font(name="맑은 고딕", size=15, bold=True, color="FFFFFF")
        section_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
        hdr_font = Font(name="맑은 고딕", size=9, bold=True, color="2D3748")
        bold_font = Font(name="맑은 고딕", size=9, bold=True)
        regular_font = Font(name="맑은 고딕", size=9)
        red_font = Font(name="맑은 고딕", size=9, bold=True, color="C53030")
        blue_font = Font(name="맑은 고딕", size=9, bold=True, color="2B6CB0")

        thin_border = Border(
            left=Side(style='thin', color='CBD5E0'), right=Side(style='thin', color='CBD5E0'),
            top=Side(style='thin', color='CBD5E0'), bottom=Side(style='thin', color='CBD5E0')
        )

        ws.merge_cells("A1:H2")
        title_cell = ws["A1"]
        title_cell.value = f"📊 [{date_str}] 키움증권 HTS 4대 시장 데이터 (0659 테마 / 0198 급등 / 0184 코스피·코스닥 거래대금 20)"
        title_cell.font, title_cell.fill, title_cell.alignment = title_font, navy_fill, Alignment(horizontal="center", vertical="center")

        current_row = 4
        ws.merge_cells(f"A{current_row}:H{current_row}")
        s1 = ws[f"A{current_row}"]
        s1.value = "🔥 [섹션 1] 0659 당일 3대 주도 테마 및 9대 핵심 대장주"
        s1.font, s1.fill, s1.alignment = section_font, theme_hdr_fill, Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[current_row].height = 26
        current_row += 1

        for col_idx, h in enumerate(["테마구분", "테마명", "평균등락률", "총거래대금(억)", "대장순위", "종목명(코드)", "현재가", "등락률"], 1):
            c = ws.cell(row=current_row, column=col_idx, value=h)
            c.font, c.fill, c.alignment, c.border = hdr_font, sub_hdr_fill, Alignment(horizontal="center", vertical="center"), thin_border
        ws.row_dimensions[current_row].height = 20
        current_row += 1

        for t_idx, t in enumerate(market_data["themes_0659"], 1):
            theme_start_row = current_row
            for s_idx, s in enumerate(t["top3_stocks"], 1):
                c_rank = ws.cell(row=current_row, column=5, value=f"{s_idx}등 대장")
                c_name = ws.cell(row=current_row, column=6, value=f"{s['name']} ({s['code']})")
                c_price = ws.cell(row=current_row, column=7, value=s['price'])
                c_rate = ws.cell(row=current_row, column=8, value=s['rate'] / 100.0)

                for c in [c_rank, c_name, c_price, c_rate]:
                    c.font, c.border = regular_font, thin_border
                c_rank.alignment = Alignment(horizontal="center", vertical="center")
                c_name.alignment = Alignment(horizontal="left", vertical="center")
                c_price.alignment, c_price.number_format = Alignment(horizontal="right", vertical="center"), '#,##0'
                c_rate.alignment, c_rate.number_format, c_rate.font = Alignment(horizontal="right", vertical="center"), '+0.00%;-0.00%;0.00%', (red_font if s['rate'] > 0 else blue_font)
                ws.row_dimensions[current_row].height = 19
                current_row += 1

            theme_end_row = current_row - 1
            for col_idx, val in [(1, f"{t_idx}위 테마"), (2, t["theme_name"]), (3, t["avg_rate"] / 100.0), (4, t["total_trade_value"])]:
                ws.merge_cells(start_row=theme_start_row, start_column=col_idx, end_row=theme_end_row, end_column=col_idx)
                c = ws.cell(row=theme_start_row, column=col_idx, value=val)
                c.font, c.alignment = bold_font, Alignment(horizontal="center", vertical="center")
                if col_idx == 3: c.number_format = '+0.00%;-0.00%;0.00%'
                elif col_idx == 4: c.number_format = '#,##0'

        current_row += 1
        ws.merge_cells(f"A{current_row}:H{current_row}")
        s2 = ws[f"A{current_row}"]
        s2.value = "⚡ [섹션 2] 0198 당일 실시간 조회수 급등 20종목 (시장 관심도 집중)"
        s2.font, s2.fill, s2.alignment = section_font, amber_hdr_fill, Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[current_row].height = 26
        current_row += 1

        for col_idx, h in enumerate(["순위", "종목코드", "종목명", "시장", "현재가(원)", "등락률(%)", "거래대금(억원)", "검색급등률"], 1):
            c = ws.cell(row=current_row, column=col_idx, value=h)
            c.font, c.fill, c.alignment, c.border = hdr_font, sub_hdr_fill, Alignment(horizontal="center", vertical="center"), thin_border
        ws.row_dimensions[current_row].height = 20
        current_row += 1

        for s in market_data["momentum_0198"]:
            row_vals = [s["rank"], s["code"], s["name"], s["market"], s["price"], s["rate"] / 100.0, s["trade_value_eok"], s["search_surge_rate"]]
            for col_idx, val in enumerate(row_vals, 1):
                c = ws.cell(row=current_row, column=col_idx, value=val)
                c.font, c.border = regular_font, thin_border
                if col_idx in [1, 2, 4, 8]: c.alignment = Alignment(horizontal="center", vertical="center")
                elif col_idx == 3: c.alignment, c.font = Alignment(horizontal="left", vertical="center"), bold_font
                elif col_idx == 5: c.alignment, c.number_format = Alignment(horizontal="right", vertical="center"), '#,##0'
                elif col_idx == 6: c.alignment, c.number_format, c.font = Alignment(horizontal="right", vertical="center"), '+0.00%;-0.00%;0.00%', (red_font if s["rate"] > 0 else blue_font)
                elif col_idx == 7: c.alignment, c.number_format = Alignment(horizontal="right", vertical="center"), '#,##0.0'
            ws.row_dimensions[current_row].height = 19
            current_row += 1

        current_row += 1
        ws.merge_cells(f"A{current_row}:H{current_row}")
        s3 = ws[f"A{current_row}"]
        s3.value = "💰 [섹션 3] 0184 KOSPI 거래대금 상위 20 순수 종목 (ETF/스팩/우선주 철저 제외)"
        s3.font, s3.fill, s3.alignment = section_font, emerald_hdr_fill, Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[current_row].height = 26
        current_row += 1

        for col_idx, h in enumerate(["순위", "종목코드", "종목명", "시장", "현재가(원)", "등락률(%)", "당일 거래대금(억원)"], 1):
            c = ws.cell(row=current_row, column=col_idx, value=h)
            c.font, c.fill, c.alignment, c.border = hdr_font, sub_hdr_fill, Alignment(horizontal="center", vertical="center"), thin_border
        ws.row_dimensions[current_row].height = 20
        current_row += 1

        for s in market_data["kospi_0184"]:
            row_vals = [s["rank"], s["code"], s["name"], s["market"], s["price"], s["rate"] / 100.0, s["trade_value_eok"]]
            for col_idx, val in enumerate(row_vals, 1):
                c = ws.cell(row=current_row, column=col_idx, value=val)
                c.font, c.border = regular_font, thin_border
                if col_idx in [1, 2, 4]: c.alignment = Alignment(horizontal="center", vertical="center")
                elif col_idx == 3: c.alignment, c.font = Alignment(horizontal="left", vertical="center"), bold_font
                elif col_idx == 5: c.alignment, c.number_format = Alignment(horizontal="right", vertical="center"), '#,##0'
                elif col_idx == 6: c.alignment, c.number_format, c.font = Alignment(horizontal="right", vertical="center"), '+0.00%;-0.00%;0.00%', (red_font if s["rate"] > 0 else blue_font)
                elif col_idx == 7: c.alignment, c.number_format = Alignment(horizontal="right", vertical="center"), '#,##0.0'
            ws.row_dimensions[current_row].height = 19
            current_row += 1

        current_row += 1
        ws.merge_cells(f"A{current_row}:H{current_row}")
        s4 = ws[f"A{current_row}"]
        s4.value = "💎 [섹션 4] 0184 KOSDAQ 거래대금 상위 20 순수 종목 (ETF/스팩/우선주 철저 제외)"
        s4.font, s4.fill, s4.alignment = section_font, blue_hdr_fill, Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[current_row].height = 26
        current_row += 1

        for col_idx, h in enumerate(["순위", "종목코드", "종목명", "시장", "현재가(원)", "등락률(%)", "당일 거래대금(억원)"], 1):
            c = ws.cell(row=current_row, column=col_idx, value=h)
            c.font, c.fill, c.alignment, c.border = hdr_font, sub_hdr_fill, Alignment(horizontal="center", vertical="center"), thin_border
        ws.row_dimensions[current_row].height = 20
        current_row += 1

        for s in market_data["kosdaq_0184"]:
            row_vals = [s["rank"], s["code"], s["name"], s["market"], s["price"], s["rate"] / 100.0, s["trade_value_eok"]]
            for col_idx, val in enumerate(row_vals, 1):
                c = ws.cell(row=current_row, column=col_idx, value=val)
                c.font, c.border = regular_font, thin_border
                if col_idx in [1, 2, 4]: c.alignment = Alignment(horizontal="center", vertical="center")
                elif col_idx == 3: c.alignment, c.font = Alignment(horizontal="left", vertical="center"), bold_font
                elif col_idx == 5: c.alignment, c.number_format = Alignment(horizontal="right", vertical="center"), '#,##0'
                elif col_idx == 6: c.alignment, c.number_format, c.font = Alignment(horizontal="right", vertical="center"), '+0.00%;-0.00%;0.00%', (red_font if s["rate"] > 0 else blue_font)
                elif col_idx == 7: c.alignment, c.number_format = Alignment(horizontal="right", vertical="center"), '#,##0.0'
            ws.row_dimensions[current_row].height = 19
            current_row += 1

        for idx, width in enumerate([12, 12, 22, 12, 14, 14, 18, 40], 1):
            ws.column_dimensions[get_column_letter(idx)].width = width

        wb.save(excel_path)
        print(f">> [ThemeCalendarWorker] ✅ 1달 단위 통합 엑셀 저장 완료: {excel_path} (시트: {date_str})")

    def _load_history(self) -> Dict[str, Any]:
        if os.path.exists(self.history_file):
            try:
                with open(self.history_file, "r", encoding="utf-8") as f: return json.load(f)
            except Exception: return {}
        return {}

    def _save_history(self, history: Dict[str, Any]):
        with open(self.history_file, "w", encoding="utf-8") as f:
            json.dump(history, f, ensure_ascii=False, indent=2)

    def _load_pullback_top30(self) -> List[Dict[str, Any]]:
        """마스터 엑셀에서 전수 누적 승률 랭킹 TOP 30 종목을 동적으로 로드"""
        candidate_paths = [
            os.path.join(self.verification_dir, "매일포착종목_통합추적관찰_마스터.xlsx"),
            r"G:\내 드라이브\Antigravity\사용자 검증\매일포착종목_통합추적관찰_마스터.xlsx"
        ]
        target_path = None
        for p in candidate_paths:
            if os.path.exists(p):
                target_path = p
                break

        if not target_path:
            return []

        try:
            wb = openpyxl.load_workbook(target_path, data_only=True)
            if "01_누적통합_승률랭킹_TOP30" not in wb.sheetnames:
                return []
            ws = wb["01_누적통합_승률랭킹_TOP30"]
            items = []
            for r in range(4, ws.max_row + 1):
                rank = ws.cell(row=r, column=1).value
                if not rank:
                    continue
                date_str = str(ws.cell(row=r, column=2).value or "")
                name = str(ws.cell(row=r, column=3).value or "")
                theme = str(ws.cell(row=r, column=4).value or "")
                win_rate = ws.cell(row=r, column=5).value or 0
                main_force = str(ws.cell(row=r, column=6).value or "")
                curr_price = ws.cell(row=r, column=7).value or 0
                entry_price = ws.cell(row=r, column=8).value or 0
                p1_price = ws.cell(row=r, column=9).value or 0
                p1_diff = ws.cell(row=r, column=10).value or 0.0
                p2_price = ws.cell(row=r, column=11).value or 0
                p2_diff = ws.cell(row=r, column=12).value or 0.0
                stop_price = ws.cell(row=r, column=13).value or 0
                status = str(ws.cell(row=r, column=14).value or "")

                items.append({
                    "rank": str(rank),
                    "capture_date": date_str,
                    "name": name,
                    "theme": theme,
                    "win_rate": int(win_rate) if isinstance(win_rate, (int, float)) else 0,
                    "main_force": main_force,
                    "curr_price": int(curr_price) if isinstance(curr_price, (int, float)) else 0,
                    "entry_price": int(entry_price) if isinstance(entry_price, (int, float)) else 0,
                    "p1_price": int(p1_price) if isinstance(p1_price, (int, float)) else 0,
                    "p1_diff": float(p1_diff) if isinstance(p1_diff, (int, float)) else 0.0,
                    "p2_price": int(p2_price) if isinstance(p2_price, (int, float)) else 0,
                    "p2_diff": float(p2_diff) if isinstance(p2_diff, (int, float)) else 0.0,
                    "stop_price": int(stop_price) if isinstance(stop_price, (int, float)) else 0,
                    "status": status
                })
            return items
        except Exception as e:
            print(f">> [ThemeCalendarWorker] ⚠️ 마스터 TOP30 로드 중 예외: {e}")
            return []

    def _build_desktop_calendar(self, history: Dict[str, Any], year: int, month: int):
        top30_data = self._load_pullback_top30()
        html = self._render_calendar_html(history, year, month, top30_data)
        
        # 1. 바탕화면 쓰기
        try:
            with open(self.desktop_html_path, "w", encoding="utf-8") as f:
                f.write(html)
            print(f">> [ThemeCalendarWorker] 📅 바탕화면 캘린더 생성 완료: {self.desktop_html_path}")
        except Exception as e:
            print(f">> [ThemeCalendarWorker] ⚠️ 바탕화면 캘린더 저장 실패: {e}")

        # 2. 구글 드라이브 캘린더 폴더 동시 저장
        try:
            gdrive_dir = os.path.dirname(self.gdrive_html_path)
            if os.path.exists(gdrive_dir):
                with open(self.gdrive_html_path, "w", encoding="utf-8") as f:
                    f.write(html)
                print(f">> [ThemeCalendarWorker] ☁️ 구글 드라이브 캘린더 동시 저장 완료: {self.gdrive_html_path}")
        except Exception as e:
            print(f">> [ThemeCalendarWorker] ⚠️ 구글 드라이브 캘린더 저장 실패: {e}")

    def _render_calendar_html(self, history: Dict[str, Any], year: int, month: int, top30_data: Optional[List[Dict[str, Any]]] = None) -> str:
        if top30_data is None:
            top30_data = self._load_pullback_top30()

        history_json_str = json.dumps(history, ensure_ascii=False)
        top30_json_str = json.dumps(top30_data, ensure_ascii=False)
        top30_count = len(top30_data)

        return f"""<!DOCTYPE html>
<html lang="ko" class="dark">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>증시 주도 테마 & HTS 4종 주도주 캘린더 (눌림목 타점 감시 통합)</title>
  <script src="https://www.gstatic.com/antigravity/web/dev/tailwindcss.min.js"></script>
  <style>
    @import url('https://fonts.googleapis.com/css2?family=Pretendard:wght@400;600;700;800&display=swap');
    body {{
      font-family: 'Pretendard', -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
    }}
    .custom-scrollbar::-webkit-scrollbar {{ width: 6px; height: 6px; }}
    .custom-scrollbar::-webkit-scrollbar-track {{ background: rgba(30, 41, 59, 0.5); }}
    .custom-scrollbar::-webkit-scrollbar-thumb {{ background: rgba(100, 116, 139, 0.5); border-radius: 3px; }}
    @keyframes pulse-slow {{
      0%, 100% {{ opacity: 1; }}
      50% {{ opacity: 0.7; }}
    }}
    .animate-pulse-slow {{
      animation: pulse-slow 2.5s cubic-bezier(0.4, 0, 0.6, 1) infinite;
    }}
  </style>
</head>
<body class="bg-slate-950 text-slate-100 min-h-screen p-4 md:p-8 antialiased selection:bg-purple-600 selection:text-white">
  
  <!-- 상단 메인 헤더 바 -->
  <div class="max-w-7xl mx-auto mb-6 flex flex-col md:flex-row md:items-center justify-between gap-4 bg-slate-900/90 border border-slate-800 p-6 rounded-2xl shadow-xl backdrop-blur">
    <div>
      <div class="flex items-center gap-3 mb-1">
        <span class="inline-flex items-center justify-center p-2 bg-purple-600/20 text-purple-400 rounded-xl border border-purple-500/30">
          <svg class="w-6 h-6" fill="none" stroke="currentColor" viewBox="0 0 24 24"><path stroke-linecap="round" stroke-linejoin="round" stroke-width="2" d="M8 7V3m8 4V3m-9 8h10M5 21h14a2 2 0 002-2V7a2 2 0 00-2-2H5a2 2 0 00-2 2v12a2 2 0 002 2z"></path></svg>
        </span>
        <h1 class="text-2xl md:text-3xl font-extrabold tracking-tight bg-clip-text text-transparent bg-gradient-to-r from-purple-400 via-indigo-300 to-blue-400">
          대한민국 증시 주도 테마 & 윗꼬리 눌림목 캘린더
        </h1>
      </div>
      <p class="text-slate-400 text-sm">
        키움증권 HTS <span class="text-blue-400 font-semibold">[0659 테마 3x3]</span> · <span class="text-amber-400 font-semibold">[0198 조회수 급등 20]</span> · <span class="text-emerald-400 font-semibold">[0184 코스피/코스닥]</span> · <span class="text-purple-400 font-bold bg-purple-950/60 px-2 py-0.5 rounded border border-purple-800/60">🎯 [눌림목 TOP 30 영구 감시]</span>
      </p>
    </div>

    <div class="flex flex-wrap items-center gap-3">
      <!-- 🎯 영구 신설 퀵 액션 메뉴 버튼 -->
      <button onclick="openPullbackModal()" class="flex items-center gap-2 px-4 py-2 bg-gradient-to-r from-purple-600 via-indigo-600 to-blue-600 hover:from-purple-500 hover:to-blue-500 text-white text-xs md:text-sm font-bold rounded-xl shadow-lg shadow-purple-500/25 border border-purple-400/30 transition transform hover:scale-105 active:scale-95 animate-pulse-slow">
        <span class="text-base">🎯</span>
        <span>윗꼬리 눌림목 TOP 30</span>
        <span class="px-2 py-0.5 bg-white/20 rounded-full text-[11px] font-extrabold">{top30_count}종목</span>
      </button>

      <div class="flex items-center gap-1.5 bg-slate-800/80 p-1.5 rounded-xl border border-slate-700">
        <button onclick="changeMonth(-1)" class="p-1.5 hover:bg-slate-700 text-slate-300 hover:text-white rounded-lg transition" title="이전 달">
          <svg class="w-5 h-5" fill="none" stroke="currentColor" viewBox="0 0 24 24"><path stroke-linecap="round" stroke-linejoin="round" stroke-width="2" d="M15 19l-7-7 7-7"></path></svg>
        </button>
        <span id="currentMonthYear" class="text-base font-bold px-2 text-slate-100 min-w-[110px] text-center"></span>
        <button onclick="changeMonth(1)" class="p-1.5 hover:bg-slate-700 text-slate-300 hover:text-white rounded-lg transition" title="다음 달">
          <svg class="w-5 h-5" fill="none" stroke="currentColor" viewBox="0 0 24 24"><path stroke-linecap="round" stroke-linejoin="round" stroke-width="2" d="M9 5l7 7-7 7"></path></svg>
        </button>
        <button onclick="goToToday()" class="px-2.5 py-1 bg-blue-600 hover:bg-blue-500 text-white text-xs font-semibold rounded-lg shadow-sm transition">오늘</button>
      </div>
    </div>
  </div>

  <!-- 요일 헤더 -->
  <div class="max-w-7xl mx-auto grid grid-cols-7 gap-2 mb-2 text-center font-bold text-sm">
    <div class="p-2 text-red-400 bg-red-950/30 rounded-lg border border-red-900/30">일 (SUN)</div>
    <div class="p-2 text-slate-300 bg-slate-900/50 rounded-lg border border-slate-800">월 (MON)</div>
    <div class="p-2 text-slate-300 bg-slate-900/50 rounded-lg border border-slate-800">화 (TUE)</div>
    <div class="p-2 text-slate-300 bg-slate-900/50 rounded-lg border border-slate-800">수 (WED)</div>
    <div class="p-2 text-slate-300 bg-slate-900/50 rounded-lg border border-slate-800">목 (THU)</div>
    <div class="p-2 text-slate-300 bg-slate-900/50 rounded-lg border border-slate-800">금 (FRI)</div>
    <div class="p-2 text-blue-400 bg-blue-950/30 rounded-lg border border-blue-900/30">토 (SAT)</div>
  </div>

  <!-- 메인 캘린더 그리드 -->
  <div id="calendarGrid" class="max-w-7xl mx-auto grid grid-cols-7 gap-2 mb-8"></div>

  <!-- 통합 상세 모달창 -->
  <div id="detailModal" class="fixed inset-0 bg-black/85 backdrop-blur-md z-50 flex items-center justify-center p-3 md:p-6 hidden">
    <div class="bg-slate-900 border border-slate-700 w-full max-w-6xl max-h-[92vh] rounded-2xl shadow-2xl flex flex-col overflow-hidden">
      
      <!-- 모달 상단 헤더 -->
      <div class="p-4 md:p-5 border-b border-slate-800 flex items-center justify-between bg-slate-950/80">
        <div class="flex items-center gap-3">
          <span id="modalIconBox" class="p-2 bg-indigo-500/20 text-indigo-400 rounded-lg border border-indigo-500/30">
            <svg class="w-5 h-5" fill="none" stroke="currentColor" viewBox="0 0 24 24"><path stroke-linecap="round" stroke-linejoin="round" stroke-width="2" d="M9 19v-6a2 2 0 00-2-2H5a2 2 0 00-2 2v6a2 2 0 002 2h2a2 2 0 002-2zm0 0V9a2 2 0 012-2h2a2 2 0 012 2v10m-6 0a2 2 0 002 2h2a2 2 0 002-2m0 0V5a2 2 0 012-2h2a2 2 0 012 2v14a2 2 0 01-2 2h-2a2 2 0 01-2-2z"></path></svg>
          </span>
          <div>
            <h2 id="modalDateTitle" class="text-lg md:text-xl font-bold text-white"></h2>
            <p id="modalSubTitle" class="text-xs text-slate-400">키움증권 HTS [0659/0198/0184] 및 전수 누적 윗꼬리 눌림목 감시 데이터</p>
          </div>
        </div>
        <button onclick="closeModal()" class="p-2 text-slate-400 hover:text-white hover:bg-slate-800 rounded-lg transition">
          <svg class="w-6 h-6" fill="none" stroke="currentColor" viewBox="0 0 24 24"><path stroke-linecap="round" stroke-linejoin="round" stroke-width="2" d="M6 18L18 6M6 6l12 12"></path></svg>
        </button>
      </div>

      <!-- 모달 탭 바 (5대 탭) -->
      <div class="flex border-b border-slate-800 bg-slate-900/90 px-4 md:px-6 gap-2 pt-2 overflow-x-auto custom-scrollbar">
        <!-- 🎯 신규 5번째 탭: 눌림목 TOP 30 (영구 보존) -->
        <button onclick="switchTab('tab_pullback')" id="btn_tab_pullback" class="px-4 py-2.5 text-xs md:text-sm font-bold border-b-2 border-purple-500 text-purple-400 transition flex items-center gap-2 whitespace-nowrap bg-purple-950/30 rounded-t-lg">
          <span>🎯 [눌림목 TOP 30] 윗꼬리 승률 & 타점 감시</span>
          <span class="px-2 py-0.5 bg-purple-900/80 text-purple-200 border border-purple-600/50 rounded-full text-[10px] font-extrabold">{top30_count}개</span>
        </button>
        <button onclick="switchTab('tab_0659')" id="btn_tab_0659" class="px-4 py-2.5 text-xs md:text-sm font-semibold border-b-2 border-transparent text-slate-400 hover:text-slate-200 transition whitespace-nowrap">
          [0659] 3대 주도 테마 & 9대 대장주
        </button>
        <button onclick="switchTab('tab_0198')" id="btn_tab_0198" class="px-4 py-2.5 text-xs md:text-sm font-semibold border-b-2 border-transparent text-slate-400 hover:text-slate-200 transition whitespace-nowrap">
          [0198] 실시간 조회수 급등 20
        </button>
        <button onclick="switchTab('tab_kospi')" id="btn_tab_kospi" class="px-4 py-2.5 text-xs md:text-sm font-semibold border-b-2 border-transparent text-slate-400 hover:text-slate-200 transition whitespace-nowrap">
          [0184] KOSPI 거래대금 TOP 20
        </button>
        <button onclick="switchTab('tab_kosdaq')" id="btn_tab_kosdaq" class="px-4 py-2.5 text-xs md:text-sm font-semibold border-b-2 border-transparent text-slate-400 hover:text-slate-200 transition whitespace-nowrap">
          [0184] KOSDAQ 거래대금 TOP 20
        </button>
      </div>

      <!-- 모달 컨텐츠 바디 -->
      <div class="p-4 md:p-6 overflow-y-auto custom-scrollbar flex-1 space-y-6">
        
        <!-- 🎯 [탭 5 컨텐츠: 눌림목 TOP 30] -->
        <div id="content_tab_pullback" class="space-y-4">
          <!-- 상단 4대 요약 카드 -->
          <div class="grid grid-cols-2 md:grid-cols-4 gap-3">
            <div class="p-3 bg-purple-950/40 border border-purple-900/60 rounded-xl">
              <div class="text-[11px] text-purple-300 font-medium">전체 감시 종목</div>
              <div class="text-xl font-extrabold text-purple-200 mt-1" id="stat_total_count">{top30_count}개</div>
              <div class="text-[10px] text-slate-400">과거+신규 누적 통합</div>
            </div>
            <div class="p-3 bg-amber-950/40 border border-amber-900/60 rounded-xl">
              <div class="text-[11px] text-amber-300 font-medium">⚡ 1차 눌림 타점권 (-1%)</div>
              <div class="text-xl font-extrabold text-amber-200 mt-1" id="stat_p1_count">0개</div>
              <div class="text-[10px] text-slate-400">시초가 대비 -1% 안착</div>
            </div>
            <div class="p-3 bg-emerald-950/40 border border-emerald-900/60 rounded-xl">
              <div class="text-[11px] text-emerald-300 font-medium">🎯 2차 골든 중심타점</div>
              <div class="text-xl font-extrabold text-emerald-200 mt-1" id="stat_p2_count">0개</div>
              <div class="text-[10px] text-slate-400">몸통 50% 중심값 매집</div>
            </div>
            <div class="p-3 bg-blue-950/40 border border-blue-900/60 rounded-xl">
              <div class="text-[11px] text-blue-300 font-medium">🚀 돌파지속 / ❌ 손절선</div>
              <div class="text-xl font-extrabold text-blue-200 mt-1" id="stat_other_count">0개</div>
              <div class="text-[10px] text-slate-400">상방 슈팅 및 리스크 관리</div>
            </div>
          </div>

          <!-- 실시간 검색 및 필터 바 -->
          <div class="flex flex-col sm:flex-row items-center justify-between gap-3 bg-slate-950/60 p-3 rounded-xl border border-slate-800">
            <div class="flex items-center gap-2 w-full sm:w-auto">
              <input type="text" id="pullbackSearchInput" onkeyup="filterPullbackTable()" placeholder="종목명 또는 테마 검색..." class="bg-slate-900 border border-slate-700 rounded-lg px-3 py-1.5 text-xs text-slate-200 focus:outline-none focus:border-purple-500 w-full sm:w-60">
              <span class="text-xs text-slate-500">실시간 검색</span>
            </div>
            <div class="flex items-center gap-1.5 flex-wrap w-full sm:w-auto justify-start sm:justify-end text-xs">
              <button onclick="setPullbackFilter('all')" class="px-2.5 py-1 rounded bg-slate-800 text-slate-200 hover:bg-slate-700 font-semibold transition text-xs filter-btn active" data-filter="all">전체보기</button>
              <button onclick="setPullbackFilter('p1')" class="px-2.5 py-1 rounded bg-slate-900 text-amber-400 hover:bg-amber-950/50 border border-amber-800/40 font-semibold transition text-xs filter-btn" data-filter="p1">⚡ 1차타점</button>
              <button onclick="setPullbackFilter('p2')" class="px-2.5 py-1 rounded bg-slate-900 text-emerald-400 hover:bg-emerald-950/50 border border-emerald-800/40 font-semibold transition text-xs filter-btn" data-filter="p2">🎯 2차타점</button>
              <button onclick="setPullbackFilter('breakout')" class="px-2.5 py-1 rounded bg-slate-900 text-blue-400 hover:bg-blue-950/50 border border-blue-800/40 font-semibold transition text-xs filter-btn" data-filter="breakout">🚀 돌파</button>
              <button onclick="setPullbackFilter('stop')" class="px-2.5 py-1 rounded bg-slate-900 text-red-400 hover:bg-red-950/50 border border-red-800/40 font-semibold transition text-xs filter-btn" data-filter="stop">❌ 손절이탈</button>
            </div>
          </div>

          <!-- TOP 30 테이블 -->
          <div class="overflow-x-auto rounded-xl border border-slate-800">
            <table class="w-full text-left text-xs md:text-sm text-slate-300">
              <thead class="bg-slate-950 text-slate-400 font-bold uppercase text-[11px] sticky top-0 z-10">
                <tr>
                  <th class="p-2.5 text-center">순위</th>
                  <th class="p-2.5 text-center">포착일</th>
                  <th class="p-2.5">종목명 (테마)</th>
                  <th class="p-2.5 text-center">예측승률</th>
                  <th class="p-2.5 text-center">수급 주포</th>
                  <th class="p-2.5 text-right">현재가</th>
                  <th class="p-2.5 text-right">1차눌림(-1%)</th>
                  <th class="p-2.5 text-right">2차중심타점</th>
                  <th class="p-2.5 text-right">절대손절선</th>
                  <th class="p-2.5 text-center">눌림목 진입 상태</th>
                </tr>
              </thead>
              <tbody id="table_pullback_body" class="divide-y divide-slate-800/60"></tbody>
            </table>
          </div>
        </div>

        <!-- 0659 테마 탭 -->
        <div id="content_tab_0659" class="space-y-6 hidden">
          <div id="themeCardsContainer" class="space-y-4"></div>
        </div>

        <!-- 0198 조회수 급등 탭 -->
        <div id="content_tab_0198" class="hidden">
          <div class="overflow-x-auto rounded-xl border border-slate-800">
            <table class="w-full text-left text-sm text-slate-300">
              <thead class="bg-slate-950 text-slate-400 font-bold uppercase text-xs">
                <tr>
                  <th class="p-3 text-center">순위</th>
                  <th class="p-3">종목명 (코드)</th>
                  <th class="p-3 text-center">시장</th>
                  <th class="p-3 text-right">현재가</th>
                  <th class="p-3 text-right">등락률</th>
                  <th class="p-3 text-right">거래대금</th>
                  <th class="p-3 text-center">검색급등률</th>
                  <th class="p-3">주도 및 급등 사유</th>
                </tr>
              </thead>
              <tbody id="table_0198_body" class="divide-y divide-slate-800/60"></tbody>
            </table>
          </div>
        </div>

        <!-- KOSPI 거래대금 탭 -->
        <div id="content_tab_kospi" class="hidden">
          <div class="mb-3 text-xs text-slate-400 flex items-center gap-1.5">
            <span class="inline-block w-2 h-2 rounded-full bg-emerald-500"></span>
            엄격한 KOSPI 순수 기업 개별주만 선별 (ETF, ETN, 우선주 제외)
          </div>
          <div class="overflow-x-auto rounded-xl border border-slate-800">
            <table class="w-full text-left text-sm text-slate-300">
              <thead class="bg-slate-950 text-slate-400 font-bold uppercase text-xs">
                <tr>
                  <th class="p-3 text-center">순위</th>
                  <th class="p-3">종목명 (코드)</th>
                  <th class="p-3 text-center">시장</th>
                  <th class="p-3 text-right">현재가</th>
                  <th class="p-3 text-right">등락률</th>
                  <th class="p-3 text-right">당일 거래대금(억원)</th>
                </tr>
              </thead>
              <tbody id="table_kospi_body" class="divide-y divide-slate-800/60"></tbody>
            </table>
          </div>
        </div>

        <!-- KOSDAQ 거래대금 탭 -->
        <div id="content_tab_kosdaq" class="hidden">
          <div class="mb-3 text-xs text-slate-400 flex items-center gap-1.5">
            <span class="inline-block w-2 h-2 rounded-full bg-blue-500"></span>
            엄격한 KOSDAQ 순수 기업 개별주만 선별 (ETF, 스팩 제외)
          </div>
          <div class="overflow-x-auto rounded-xl border border-slate-800">
            <table class="w-full text-left text-sm text-slate-300">
              <thead class="bg-slate-950 text-slate-400 font-bold uppercase text-xs">
                <tr>
                  <th class="p-3 text-center">순위</th>
                  <th class="p-3">종목명 (코드)</th>
                  <th class="p-3 text-center">시장</th>
                  <th class="p-3 text-right">현재가</th>
                  <th class="p-3 text-right">등락률</th>
                  <th class="p-3 text-right">당일 거래대금(억원)</th>
                </tr>
              </thead>
              <tbody id="table_kosdaq_body" class="divide-y divide-slate-800/60"></tbody>
            </table>
          </div>
        </div>

      </div>
    </div>
  </div>

  <script>
    const marketHistory = {history_json_str};
    const pullbackTop30 = {top30_json_str};
    let currentYear = {year};
    let currentMonth = {month};
    let currentPullbackFilter = 'all';

    function renderCalendar() {{
      const grid = document.getElementById("calendarGrid");
      grid.innerHTML = "";

      document.getElementById("currentMonthYear").innerText = `${{currentYear}}년 ${{currentMonth}}월`;

      const firstDayIndex = new Date(currentYear, currentMonth - 1, 1).getDay();
      const lastDate = new Date(currentYear, currentMonth, 0).getDate();
      const prevMonthLastDate = new Date(currentYear, currentMonth - 1, 0).getDate();

      for (let i = firstDayIndex - 1; i >= 0; i--) {{
        const day = prevMonthLastDate - i;
        const cell = document.createElement("div");
        cell.className = "min-h-[140px] p-2.5 bg-slate-950/40 border border-slate-900 rounded-xl opacity-30";
        cell.innerHTML = `<span class="text-xs font-semibold text-slate-600">${{day}}</span>`;
        grid.appendChild(cell);
      }}

      for (let day = 1; day <= lastDate; day++) {{
        const dateStr = `${{currentYear}}-${{String(currentMonth).padStart(2, '0')}}-${{String(day).padStart(2, '0')}}`;
        const dayData = marketHistory[dateStr];
        const dayOfWeek = new Date(currentYear, currentMonth - 1, day).getDay();

        const isToday = (dateStr === new Date().toISOString().slice(0, 10));
        const isWeekend = (dayOfWeek === 0 || dayOfWeek === 6);

        const cell = document.createElement("div");
        let borderClass = isToday ? "border-purple-500 shadow-lg shadow-purple-500/20 ring-1 ring-purple-500" : (isWeekend ? "border-slate-900 opacity-60" : "border-slate-800 hover:border-slate-700");
        let bgClass = isToday ? "bg-slate-900/90" : (isWeekend ? "bg-slate-950/60" : "bg-slate-900/50 hover:bg-slate-900/80");

        cell.className = `min-h-[145px] p-2.5 rounded-xl border ${{borderClass}} ${{bgClass}} transition ${{isWeekend ? 'cursor-default' : 'cursor-pointer'}} flex flex-col justify-between group`;
        if (!isWeekend) {{
          cell.onclick = () => openModal(dateStr);
        }}

        let themesHtml = "";
        if (dayData && dayData.themes_0659 && dayData.themes_0659.length > 0) {{
          themesHtml = `<div class="mt-2 space-y-1">`;
          dayData.themes_0659.slice(0, 3).forEach((t, idx) => {{
            const badgeColor = idx === 0 ? "bg-amber-500/20 text-amber-400 border border-amber-500/30" : (idx === 1 ? "bg-slate-700/40 text-slate-300" : "bg-orange-500/20 text-orange-400");
            themesHtml += `
              <div class="text-[11px] font-medium truncate py-0.5 px-1.5 rounded ${{badgeColor}} flex items-center justify-between">
                <span class="truncate">${{idx + 1}}. ${{t.theme_name}}</span>
                <span class="font-bold text-[10px] text-red-400">+${{t.avg_rate.toFixed(1)}}%</span>
              </div>
            `;
          }});
          themesHtml += `</div>`;
        }} else if (!isWeekend) {{
          themesHtml = `<div class="mt-2 text-xs text-slate-600 italic">데이터 없음</div>`;
        }}

        // 포착 종목 알림 배지 (오늘 또는 포착 데이터 존재 시)
        let pullbackBadge = "";
        const capturedToday = pullbackTop30.filter(p => p.capture_date === dateStr);
        if (capturedToday.length > 0 || isToday) {{
          const cnt = capturedToday.length > 0 ? capturedToday.length : pullbackTop30.length;
          pullbackBadge = `
            <div onclick="event.stopPropagation(); openPullbackModal();" class="mt-1.5 px-2 py-1 bg-purple-950/80 border border-purple-800/80 rounded-lg text-purple-300 text-[10px] font-bold flex items-center justify-between hover:bg-purple-900 transition" title="클릭 시 눌림목 TOP 30 즉시 열람">
              <span>🎯 눌림목 TOP30</span>
              <span class="bg-purple-800/60 px-1 py-0.2 rounded text-[9px] text-purple-200 font-extrabold">${{cnt}}개</span>
            </div>
          `;
        }}

        cell.innerHTML = `
          <div>
            <div class="flex items-center justify-between">
              <span class="text-sm font-bold ${{isToday ? 'text-purple-400' : (dayOfWeek === 0 ? 'text-red-400' : (dayOfWeek === 6 ? 'text-blue-400' : 'text-slate-300'))}}">${{day}}</span>
              ${{isToday ? '<span class="px-1.5 py-0.5 text-[10px] font-bold bg-purple-600 text-white rounded">TODAY</span>' : ''}}
            </div>
            ${{themesHtml}}
          </div>
          <div>
            ${{pullbackBadge}}
          </div>
        `;
        grid.appendChild(cell);
      }}
    }}

    function changeMonth(delta) {{
      currentMonth += delta;
      if (currentMonth > 12) {{ currentMonth = 1; currentYear++; }}
      else if (currentMonth < 1) {{ currentMonth = 12; currentYear--; }}
      renderCalendar();
    }}

    function goToToday() {{
      const d = new Date();
      currentYear = d.getFullYear();
      currentMonth = d.getMonth() + 1;
      renderCalendar();
    }}

    // 🎯 영구 신설: 눌림목 TOP 30 모달 즉시 열기
    function openPullbackModal() {{
      const modal = document.getElementById("detailModal");
      document.getElementById("modalDateTitle").innerText = "🎯 [전수 누적 추적] 윗꼬리 매집봉 승률 TOP 30 & 실시간 눌림목 타점 감시";
      document.getElementById("modalSubTitle").innerText = "마스터 엑셀(01_누적통합_승률랭킹_TOP30) 연동 · 과거 및 신규 포착 전수 통합 추적";
      
      renderPullbackTable();
      switchTab('tab_pullback');
      modal.classList.remove("hidden");
    }}

    function openModal(dateStr) {{
      const data = marketHistory[dateStr];
      const modal = document.getElementById("detailModal");
      document.getElementById("modalDateTitle").innerText = `📅 [${{dateStr}}] 키움증권 HTS 4종 주도 시장 데이터 리포트`;
      document.getElementById("modalSubTitle").innerText = "키움증권 HTS [0659/0198/0184] 및 전수 누적 윗꼬리 눌림목 감시 데이터";

      renderPullbackTable();

      if (!data) {{
        // HTS 데이터가 없는 날이라도 눌림목 TOP 30 탭으로 기본 오픈
        switchTab('tab_pullback');
        modal.classList.remove("hidden");
        return;
      }}

      const container = document.getElementById("themeCardsContainer");
      container.innerHTML = "";
      if (data.themes_0659) {{
        data.themes_0659.forEach((t, idx) => {{
          const medal = idx === 0 ? "🥇 1위 테마" : (idx === 1 ? "🥈 2위 테마" : "🥉 3위 테마");
          const card = document.createElement("div");
          card.className = "bg-slate-950/80 border border-slate-800 rounded-xl p-4";

          let stocksTable = `
            <table class="w-full text-sm text-left mt-3">
              <thead class="text-xs text-slate-400 bg-slate-900 border-b border-slate-800">
                <tr>
                  <th class="p-2 text-center">대장순위</th>
                  <th class="p-2">종목명 (코드)</th>
                  <th class="p-2 text-center">시장</th>
                  <th class="p-2 text-right">현재가</th>
                  <th class="p-2 text-right">등락률</th>
                  <th class="p-2 text-right">거래대금</th>
                  <th class="p-2">주도 사유</th>
                </tr>
              </thead>
              <tbody class="divide-y divide-slate-800/50">
          `;

          t.top3_stocks.forEach((s, sIdx) => {{
            stocksTable += `
              <tr>
                <td class="p-2 text-center font-bold text-amber-400">${{sIdx + 1}}등 대장</td>
                <td class="p-2 font-semibold text-slate-200">${{s.name}} <span class="text-xs text-slate-500 font-normal">(${{s.code}})</span></td>
                <td class="p-2 text-center"><span class="px-1.5 py-0.5 text-[10px] font-bold rounded ${{s.market === 'KOSPI' ? 'bg-emerald-900/60 text-emerald-300' : 'bg-blue-900/60 text-blue-300'}}">${{s.market}}</span></td>
                <td class="p-2 text-right">${{s.price.toLocaleString()}}원</td>
                <td class="p-2 text-right font-bold ${{s.rate > 0 ? 'text-red-400' : 'text-blue-400'}}">${{s.rate > 0 ? '+' : ''}}${{s.rate.toFixed(2)}}%</td>
                <td class="p-2 text-right">${{s.trade_value_eok.toLocaleString()}}억원</td>
                <td class="p-2 text-xs text-slate-400">${{s.reason}}</td>
              </tr>
            `;
          }});

          stocksTable += `</tbody></table>`;

          card.innerHTML = `
            <div class="flex flex-col md:flex-row md:items-center justify-between gap-2 border-b border-slate-800/80 pb-2.5">
              <div>
                <span class="px-2 py-0.5 rounded text-xs font-bold bg-indigo-600/20 text-indigo-400 border border-indigo-500/30">${{medal}}</span>
                <h3 class="text-lg font-bold text-white mt-1">${{t.theme_name}}</h3>
                <p class="text-xs text-slate-400 italic">${{t.description}}</p>
              </div>
              <div class="text-right text-xs">
                <div class="text-slate-400">평균 등락: <span class="font-bold text-red-400">+${{t.avg_rate.toFixed(2)}}%</span></div>
                <div class="text-slate-400">총 거래대금: <span class="font-bold text-slate-200">${{t.total_trade_value.toLocaleString()}}억원</span></div>
              </div>
            </div>
            ${{stocksTable}}
          `;
          container.appendChild(card);
        }});
      }}

      const table0198 = document.getElementById("table_0198_body");
      table0198.innerHTML = "";
      if (data.momentum_0198) {{
        data.momentum_0198.forEach(s => {{
          const tr = document.createElement("tr");
          tr.innerHTML = `
            <td class="p-3 text-center font-bold text-amber-400">${{s.rank}}</td>
            <td class="p-3 font-semibold text-white">${{s.name}} <span class="text-xs text-slate-500 font-normal">(${{s.code}})</span></td>
            <td class="p-3 text-center"><span class="px-1.5 py-0.5 text-[10px] font-bold rounded ${{s.market === 'KOSPI' ? 'bg-emerald-900/60 text-emerald-300' : 'bg-blue-900/60 text-blue-300'}}">${{s.market}}</span></td>
            <td class="p-3 text-right">${{s.price.toLocaleString()}}원</td>
            <td class="p-3 text-right font-bold text-red-400">+${{s.rate.toFixed(2)}}%</td>
            <td class="p-3 text-right">${{s.trade_value_eok.toLocaleString()}}억</td>
            <td class="p-3 text-center font-bold text-amber-300">${{s.search_surge_rate}}</td>
            <td class="p-3 text-xs text-slate-400">${{s.reason || '-'}}</td>
          `;
          table0198.appendChild(tr);
        }});
      }}

      const tableKospi = document.getElementById("table_kospi_body");
      tableKospi.innerHTML = "";
      if (data.kospi_0184) {{
        data.kospi_0184.forEach(s => {{
          const tr = document.createElement("tr");
          tr.innerHTML = `
            <td class="p-3 text-center font-bold text-emerald-400">${{s.rank}}</td>
            <td class="p-3 font-semibold text-white">${{s.name}} <span class="text-xs text-slate-500 font-normal">(${{s.code}})</span></td>
            <td class="p-3 text-center"><span class="px-1.5 py-0.5 text-[10px] font-bold rounded bg-emerald-900/60 text-emerald-300">KOSPI</span></td>
            <td class="p-3 text-right">${{s.price.toLocaleString()}}원</td>
            <td class="p-3 text-right font-bold ${{s.rate >= 0 ? 'text-red-400' : 'text-blue-400'}}">${{s.rate >= 0 ? '+' : ''}}${{s.rate.toFixed(2)}}%</td>
            <td class="p-3 text-right font-bold text-slate-100">${{s.trade_value_eok.toLocaleString()}}억원</td>
          `;
          tableKospi.appendChild(tr);
        }});
      }}

      const tableKosdaq = document.getElementById("table_kosdaq_body");
      tableKosdaq.innerHTML = "";
      if (data.kosdaq_0184) {{
        data.kosdaq_0184.forEach(s => {{
          const tr = document.createElement("tr");
          tr.innerHTML = `
            <td class="p-3 text-center font-bold text-blue-400">${{s.rank}}</td>
            <td class="p-3 font-semibold text-white">${{s.name}} <span class="text-xs text-slate-500 font-normal">(${{s.code}})</span></td>
            <td class="p-3 text-center"><span class="px-1.5 py-0.5 text-[10px] font-bold rounded bg-blue-900/60 text-blue-300">KOSDAQ</span></td>
            <td class="p-3 text-right">${{s.price.toLocaleString()}}원</td>
            <td class="p-3 text-right font-bold ${{s.rate >= 0 ? 'text-red-400' : 'text-blue-400'}}">${{s.rate >= 0 ? '+' : ''}}${{s.rate.toFixed(2)}}%</td>
            <td class="p-3 text-right font-bold text-slate-100">${{s.trade_value_eok.toLocaleString()}}억원</td>
          `;
          tableKosdaq.appendChild(tr);
        }});
      }}

      switchTab('tab_0659');
      modal.classList.remove("hidden");
    }}

    function closeModal() {{
      document.getElementById("detailModal").classList.add("hidden");
    }}

    // 탭 전환 핸들러 (5개 탭 완벽 지원)
    function switchTab(tabId) {{
      const tabs = ['tab_pullback', 'tab_0659', 'tab_0198', 'tab_kospi', 'tab_kosdaq'];
      tabs.forEach(t => {{
        const btn = document.getElementById(`btn_${{t}}`);
        const content = document.getElementById(`content_${{t}}`);
        if (!btn || !content) return;

        if (t === tabId) {{
          if (t === 'tab_pullback') {{
            btn.className = "px-4 py-2.5 text-xs md:text-sm font-bold border-b-2 border-purple-500 text-purple-400 transition flex items-center gap-2 whitespace-nowrap bg-purple-950/40 rounded-t-lg";
          }} else {{
            btn.className = "px-4 py-2.5 text-xs md:text-sm font-semibold border-b-2 border-blue-500 text-blue-400 transition whitespace-nowrap bg-slate-800/50 rounded-t-lg";
          }}
          content.classList.remove("hidden");
        }} else {{
          btn.className = "px-4 py-2.5 text-xs md:text-sm font-semibold border-b-2 border-transparent text-slate-400 hover:text-slate-200 transition whitespace-nowrap";
          content.classList.add("hidden");
        }}
      }});
    }}

    // 🎯 [눌림목 TOP 30 테이블 렌더러]
    function renderPullbackTable() {{
      const tbody = document.getElementById("table_pullback_body");
      if (!tbody) return;
      tbody.innerHTML = "";

      const query = (document.getElementById("pullbackSearchInput")?.value || "").trim().toLowerCase();

      let p1Count = 0;
      let p2Count = 0;
      let otherCount = 0;

      pullbackTop30.forEach(item => {{
        if (item.status.includes("1차")) p1Count++;
        else if (item.status.includes("2차")) p2Count++;
        else otherCount++;
      }});

      if (document.getElementById("stat_p1_count")) document.getElementById("stat_p1_count").innerText = `${{p1Count}}개`;
      if (document.getElementById("stat_p2_count")) document.getElementById("stat_p2_count").innerText = `${{p2Count}}개`;
      if (document.getElementById("stat_other_count")) document.getElementById("stat_other_count").innerText = `${{otherCount}}개`;

      const filtered = pullbackTop30.filter(item => {{
        // 검색어 필터
        const matchQuery = !query || item.name.toLowerCase().includes(query) || item.theme.toLowerCase().includes(query);
        if (!matchQuery) return false;

        // 상태 버튼 필터
        if (currentPullbackFilter === 'p1') return item.status.includes("1차");
        if (currentPullbackFilter === 'p2') return item.status.includes("2차");
        if (currentPullbackFilter === 'breakout') return item.status.includes("돌파");
        if (currentPullbackFilter === 'stop') return item.status.includes("손절");
        return true;
      }});

      if (filtered.length === 0) {{
        tbody.innerHTML = `<tr><td colspan="10" class="p-8 text-center text-slate-500 italic">조건에 일치하는 종목이 없습니다.</td></tr>`;
        return;
      }}

      filtered.forEach((item, idx) => {{
        const tr = document.createElement("tr");
        tr.className = "hover:bg-slate-800/40 transition border-b border-slate-800/40";

        // 순위 배지
        let rankBadge = `<span class="px-2 py-0.5 rounded text-xs font-bold bg-slate-800 text-slate-300">${{item.rank}}</span>`;
        if (item.rank === "1위") rankBadge = `<span class="px-2 py-0.5 rounded text-xs font-extrabold bg-amber-500/20 text-amber-400 border border-amber-500/30">🥇 1위</span>`;
        else if (item.rank === "2위") rankBadge = `<span class="px-2 py-0.5 rounded text-xs font-extrabold bg-slate-400/20 text-slate-200 border border-slate-400/30">🥈 2위</span>`;
        else if (item.rank === "3위") rankBadge = `<span class="px-2 py-0.5 rounded text-xs font-extrabold bg-orange-600/20 text-orange-400 border border-orange-500/30">🥉 3위</span>`;
        else if (idx < 10) rankBadge = `<span class="px-2 py-0.5 rounded text-xs font-bold bg-purple-900/30 text-purple-300 border border-purple-800/40">${{item.rank}}</span>`;

        // 포착일자 배지
        const captureDateBadge = `<span class="px-1.5 py-0.5 rounded text-[11px] bg-slate-900 text-slate-400 border border-slate-800">${{item.capture_date.slice(5)}}</span>`;

        // 승률 배지 & 프로그레스 바
        const winRateColor = item.win_rate >= 70 ? "text-emerald-400" : (item.win_rate >= 67 ? "text-blue-400" : "text-amber-400");
        const winProgressColor = item.win_rate >= 70 ? "bg-emerald-500" : (item.win_rate >= 67 ? "bg-blue-500" : "bg-amber-500");

        // 수급주포 배지
        let forceBadge = `<span class="px-2 py-0.5 rounded text-[11px] font-bold bg-slate-800 text-slate-300">${{item.main_force}}</span>`;
        if (item.main_force.includes("쌍끌이")) forceBadge = `<span class="px-2 py-0.5 rounded text-[11px] font-extrabold bg-amber-500/20 text-amber-300 border border-amber-500/30">⚡ 쌍끌이</span>`;
        else if (item.main_force.includes("외인")) forceBadge = `<span class="px-2 py-0.5 rounded text-[11px] font-bold bg-purple-900/50 text-purple-300 border border-purple-700/40">외인 주포</span>`;
        else if (item.main_force.includes("기관")) forceBadge = `<span class="px-2 py-0.5 rounded text-[11px] font-bold bg-blue-900/50 text-blue-300 border border-blue-700/40">기관 주포</span>`;

        // 상태 배지
        let statusBadge = `<span class="px-2 py-0.5 rounded text-xs font-semibold bg-slate-800 text-slate-400">${{item.status}}</span>`;
        if (item.status.includes("1차")) statusBadge = `<span class="px-2 py-0.5 rounded text-xs font-bold bg-amber-500/20 text-amber-300 border border-amber-500/30">⚡ 1차 타점권 (-1%)</span>`;
        else if (item.status.includes("2차")) statusBadge = `<span class="px-2 py-0.5 rounded text-xs font-bold bg-emerald-500/20 text-emerald-300 border border-emerald-500/30">🎯 2차 골든 중심타점</span>`;
        else if (item.status.includes("돌파")) statusBadge = `<span class="px-2 py-0.5 rounded text-xs font-bold bg-blue-500/20 text-blue-300 border border-blue-500/30">🚀 타점 상단 (돌파)</span>`;
        else if (item.status.includes("손절")) statusBadge = `<span class="px-2 py-0.5 rounded text-xs font-bold bg-red-500/20 text-red-300 border border-red-500/30">❌ 손절선 이탈</span>`;

        // 1차/2차 괴리율 표시
        const p1DiffColor = item.p1_diff > 0 ? "text-red-400" : "text-blue-400";
        const p2DiffColor = item.p2_diff > 0 ? "text-red-400" : "text-blue-400";

        tr.innerHTML = `
          <td class="p-2.5 text-center">${{rankBadge}}</td>
          <td class="p-2.5 text-center">${{captureDateBadge}}</td>
          <td class="p-2.5">
            <div class="font-bold text-white text-sm hover:text-purple-300 cursor-pointer transition">${{item.name}}</div>
            <div class="text-[11px] text-slate-400 truncate max-w-[180px]">${{item.theme}}</div>
          </td>
          <td class="p-2.5 text-center">
            <div class="font-extrabold ${{winRateColor}} text-sm">${{item.win_rate}}%</div>
            <div class="w-16 bg-slate-800 rounded-full h-1.5 mx-auto mt-1 overflow-hidden">
              <div class="${{winProgressColor}} h-1.5 rounded-full" style="width: ${{item.win_rate}}%"></div>
            </div>
          </td>
          <td class="p-2.5 text-center">${{forceBadge}}</td>
          <td class="p-2.5 text-right font-bold text-slate-100">${{item.curr_price.toLocaleString()}}원</td>
          <td class="p-2.5 text-right">
            <div class="font-semibold text-amber-300">${{item.p1_price.toLocaleString()}}원</div>
            <div class="text-[10px] ${{p1DiffColor}}">${{item.p1_diff > 0 ? '+' : ''}}${{item.p1_diff.toFixed(2)}}%</div>
          </td>
          <td class="p-2.5 text-right">
            <div class="font-semibold text-emerald-300">${{item.p2_price.toLocaleString()}}원</div>
            <div class="text-[10px] ${{p2DiffColor}}">${{item.p2_diff > 0 ? '+' : ''}}${{item.p2_diff.toFixed(2)}}%</div>
          </td>
          <td class="p-2.5 text-right font-semibold text-red-400/90">${{item.stop_price.toLocaleString()}}원</td>
          <td class="p-2.5 text-center">${{statusBadge}}</td>
        `;
        tbody.appendChild(tr);
      }});
    }}

    function filterPullbackTable() {{
      renderPullbackTable();
    }}

    function setPullbackFilter(filterType) {{
      currentPullbackFilter = filterType;
      document.querySelectorAll(".filter-btn").forEach(btn => {{
        if (btn.dataset.filter === filterType) {{
          btn.classList.add("bg-purple-600", "text-white");
          btn.classList.remove("bg-slate-900", "bg-slate-800");
        }} else {{
          btn.classList.remove("bg-purple-600", "text-white");
          btn.classList.add("bg-slate-900");
        }}
      }});
      renderPullbackTable();
    }}

    // 초기 캘린더 렌더링
    renderCalendar();
  </script>
</body>
</html>"""

if __name__ == "__main__":
    worker = ThemeCalendarWorker()
    worker.run()
